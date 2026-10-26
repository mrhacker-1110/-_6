"""
Запуск:  python manage.py export_tests_docx [--app crud] [--output tests_report.docx]

Запускает тесты Django, собирает результаты и записывает в Word-файл:
таблицу результатов и исходный код всех тестов.
"""
import inspect
import unittest
from datetime import datetime

from django.apps import apps
from django.core.management.base import BaseCommand
from django.test.runner import DiscoverRunner
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = 'Times New Roman'
CODE_FONT = 'Courier New'
RESULTS = {}


class RecordingResult(unittest.TextTestResult):
    """Запоминает статус каждого теста."""

    def addSuccess(self, test):
        super().addSuccess(test)
        RESULTS[test.id()] = 'OK'

    def addFailure(self, test, err):
        super().addFailure(test, err)
        RESULTS[test.id()] = 'FAIL'

    def addError(self, test, err):
        super().addError(test, err)
        RESULTS[test.id()] = 'ERROR'


class RecordingRunner(DiscoverRunner):
    def get_resultclass(self):
        return RecordingResult


def set_font(run, name=FONT, size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_par(doc, text='', size=12, bold=False, align=None, font=FONT, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    set_font(p.add_run(text), font, size, bold)
    return p


def add_heading(doc, text, size=14):
    p = add_par(doc, text, size=size, bold=True, space_after=6)
    p.paragraph_format.keep_with_next = True
    return p


def set_cell(cell, text, bold=False, align=None, font=FONT, size=11):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    set_font(p.add_run(text), font, size, bold)


def set_borders(table):
    tbl = table._tbl
    props = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    props.append(borders)


def add_code_block(doc, code):
    """Блок кода: рамка, моноширинный шрифт, чёрно-белый."""
    table = doc.add_table(rows=1, cols=1)
    set_borders(table)
    cell = table.rows[0].cells[0]
    cell.text = ''
    lines = code.split('\n')
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(line if line else ' '), CODE_FONT, 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


class Command(BaseCommand):
    help = 'Запускает тесты Django и записывает код и результаты в Word-файл.'

    def add_arguments(self, parser):
        parser.add_argument('--app', default='crud', help='Имя приложения (по умолчанию crud)')
        parser.add_argument('--output', default='tests_report.docx', help='Имя выходного файла')

    def handle(self, *args, **opts):
        app_label = opts['app']
        output = opts['output']
        tests_module = apps.get_app_config(app_label).module.__name__ + '.tests'

        # 1. Запуск тестов
        RESULTS.clear()
        runner = RecordingRunner(verbosity=0, interactive=False)
        failures = runner.run_tests([app_label])

        # 2. Сбор тестов из модуля
        import importlib
        module = importlib.import_module(tests_module)
        classes = [
            c for _, c in inspect.getmembers(module, inspect.isclass)
            if issubclass(c, unittest.TestCase) and c.__module__ == tests_module
        ]
        classes.sort(key=lambda c: inspect.getsourcelines(c)[1])

        # 3. Документ
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
        sec.top_margin = sec.bottom_margin = Cm(2)

        style = doc.styles['Normal']
        style.font.name = FONT
        style.font.size = Pt(12)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

        add_par(doc, 'Отчёт о тестировании Django-проекта', 16, True, WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        add_par(doc, f'Приложение: {app_label} (CRUD)', 12, False, WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        add_par(doc, f'Дата: {datetime.now():%d.%m.%Y %H:%M}', 12, False, WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

        total = len(RESULTS)
        passed = sum(1 for v in RESULTS.values() if v == 'OK')
        add_par(doc, f'Всего тестов: {total}. Пройдено: {passed}. Не пройдено: {total - passed}.', bold=True)

        # Раздел 1: таблица результатов
        add_heading(doc, '1. Результаты тестов')
        n = 0
        for cls in classes:
            add_par(doc, f'{cls.__name__} — {inspect.getdoc(cls) or ""}', 12, True, space_after=4).paragraph_format.keep_with_next = True
            methods = [(name, f) for name, f in inspect.getmembers(cls, inspect.isfunction) if name.startswith('test')]
            methods.sort(key=lambda m: inspect.getsourcelines(m[1])[1])
            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_borders(table)
            widths = [Cm(1.0), Cm(6.4), Cm(7.0), Cm(2.1)]
            for i, h in enumerate(['№', 'Тест', 'Что проверяет', 'Итог']):
                set_cell(table.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            for name, func in methods:
                n += 1
                status = RESULTS.get(f'{cls.__module__}.{cls.__name__}.{name}', '—')
                row = table.add_row().cells
                set_cell(row[0], str(n), align=WD_ALIGN_PARAGRAPH.CENTER)
                set_cell(row[1], name, font=CODE_FONT, size=9)
                set_cell(row[2], inspect.getdoc(func) or '')
                set_cell(row[3], status, align=WD_ALIGN_PARAGRAPH.CENTER)
            table.autofit = False
            for i, w in enumerate(widths):
                table.columns[i].width = w
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = w
            add_par(doc, '', space_after=4)

        # Раздел 2: исходный код
        doc.add_page_break()
        add_heading(doc, '2. Исходный код тестов')
        add_par(doc, f'Файл: {app_label}/tests.py', space_after=4)
        add_code_block(doc, inspect.getsource(module).rstrip())

        doc.save(output)
        self.stdout.write(self.style.SUCCESS(f'Готово: {output} (тестов: {total}, успешно: {passed})'))
        if failures:
            self.stderr.write('Есть непройденные тесты!')
